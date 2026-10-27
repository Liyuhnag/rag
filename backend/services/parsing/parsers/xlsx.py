# -*- coding: utf-8 -*-
"""XLSX 解析器。"""

import io
import json
import zipfile
from collections.abc import Iterator
from pathlib import Path
from tempfile import TemporaryDirectory
from xml.etree.ElementTree import ParseError
from datetime import date
from datetime import datetime
from math import isfinite

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.worksheet import Worksheet

from services.parsing.adapters.object_reader import StorageObjectReader
from services.parsing.errors import ParseFailed
from services.parsing.errors import UnsupportedFileType
from services.parsing.models import ParseInput
from services.parsing.models import ParsedBlock
from services.parsing.models import ParsedDocument
from services.parsing.parsers.base import BaseParser
from services.parsing.parsers._xlsx_stream import iter_xlsx
from services.parsing.ports import ObjectReader

_MAX_SHEETS = 50
_MAX_ROWS = 200_000
_MAX_COLUMNS = 256
_MAX_CELLS = 2_000_000
_MAX_UNCOMPRESSED = 2 * 1024 * 1024 * 1024


class XlsxParser(BaseParser):
    """每个工作表转成表头为键的 JSON 数组。"""

    file_types = ("xlsx",)

    def __init__(
        self,
        reader: ObjectReader | None = None,
        *,
        max_sheets: int = _MAX_SHEETS,
        max_rows: int = _MAX_ROWS,
        max_columns: int = _MAX_COLUMNS,
        max_cells: int = _MAX_CELLS,
        max_uncompressed: int = _MAX_UNCOMPRESSED,
        batch_rows: int = 128,
    ) -> None:
        if batch_rows < 1:
            raise ValueError("batch_rows 必须大于 0")
        self._batch_rows = batch_rows
        self._reader = reader or StorageObjectReader()
        self._max_sheets = max_sheets
        self._max_rows = max_rows
        self._max_columns = max_columns
        self._max_cells = max_cells
        self._max_uncompressed = max_uncompressed

    def iter_blocks(self, source: ParseInput) -> Iterator[ParsedBlock]:
        """下载后借助临时 SQLite 逐批输出；调用方应关闭提前结束的迭代器。"""
        if not self.supports(source.request.file_type):
            raise UnsupportedFileType(
                f"解析器不支持文件类型: {source.request.file_type}"
            )
        with TemporaryDirectory(prefix="rag-xlsx-") as directory:
            path = Path(directory) / "source.xlsx"
            self._reader.download_object(source.request.object_key, path)
            try:
                yield from iter_xlsx(
                    path,
                    Path(directory) / "index.sqlite",
                    self,
                    _value,
                    (_headers, _extend_headers, _check_cells),
                )
            except (
                zipfile.BadZipFile,
                ParseError,
                ValueError,
                KeyError,
                IndexError,
                OSError,
                OverflowError,
            ) as exc:
                raise ParseFailed("XLSX 无法解析") from exc

    def _parse(self, source: ParseInput) -> ParsedDocument:
        raw = self._reader.read_object(source.request.object_key)
        return self.parse_bytes(raw, source)

    def parse_bytes(self, content: bytes, source: ParseInput) -> ParsedDocument:
        """解析已经读出的 XLSX 字节。"""
        _check_uncompressed(content, self._max_uncompressed)
        try:
            workbook = load_workbook(
                io.BytesIO(content),
                data_only=True,
                keep_links=False,
            )
        except (InvalidFileException, zipfile.BadZipFile, ValueError, OSError) as exc:
            raise ParseFailed("XLSX 无法解析") from exc
        try:
            blocks = _blocks(
                workbook.worksheets,
                self._max_sheets,
                self._max_rows,
                self._max_columns,
                self._max_cells,
            )
        finally:
            workbook.close()
        if not blocks:
            raise ParseFailed("Excel 没有表头")
        return ParsedDocument(
            request=source.request,
            title=_title(source.request.filename),
            blocks=blocks,
        )


def _check_uncompressed(content: bytes, limit: int) -> None:
    """拒绝解压后超过上限的工作簿。"""
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            total = 0
            for info in archive.infolist():
                total += max(info.file_size, 0)
                if total > limit:
                    raise ParseFailed(f"Excel 解压后超过 {limit} 字节")
    except zipfile.BadZipFile as exc:
        raise ParseFailed("XLSX 无法解析") from exc


def _blocks(
    worksheets: list[Worksheet],
    max_sheets: int,
    max_rows: int,
    max_columns: int,
    max_cells: int,
) -> list[ParsedBlock]:
    """按工作表生成表格块。"""
    if len(worksheets) > max_sheets:
        raise ParseFailed(f"Excel 超过 {max_sheets} 个工作表")
    blocks: list[ParsedBlock] = []
    cells = 0
    for worksheet in worksheets:
        records, column_count, cells = _sheet_records(
            worksheet,
            max_rows,
            max_columns,
            max_cells,
            cells,
        )
        if records is None:
            continue
        blocks.append(
            ParsedBlock(
                kind="table",
                text=json.dumps(records, ensure_ascii=False),
                sheet=worksheet.title,
                metadata={
                    "row_count": len(records),
                    "column_count": column_count,
                },
            )
        )
    return blocks


def _sheet_records(
    worksheet: Worksheet,
    max_rows: int,
    max_columns: int,
    max_cells: int,
    cells: int,
) -> tuple[list[dict[str, object]] | None, int, int]:
    """读取一张表。没有表头时返回 None。"""
    if (worksheet.max_column or 0) > max_columns:
        raise ParseFailed(f"Excel 超过 {max_columns} 列")
    if (worksheet.max_row or 0) > max_rows + 1:
        raise ParseFailed(f"Excel 超过 {max_rows} 行")
    fills = _merged_fills(worksheet, max_columns, max_cells)
    headers: list[str] = []
    used: set[str] = set()
    records: list[dict[str, object]] = []
    for row_number, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
        values = _materialize(list(row), row_number, fills)
        if not headers:
            if _blank(values):
                continue
            headers = _headers(values, used)
            cells += len(headers)
            _check_cells(cells, max_cells)
            continue
        if _blank(values):
            continue
        if len(records) >= max_rows:
            raise ParseFailed(f"Excel 超过 {max_rows} 行")
        if len(values) > len(headers):
            _extend_headers(headers, used, len(values), records)
        width = len(headers)
        cells += width
        _check_cells(cells, max_cells)
        normalized = [_value(cell) for cell in values]
        normalized.extend([None] * (width - len(normalized)))
        records.append(dict(zip(headers, normalized, strict=True)))
    if not headers:
        return None, 0, cells
    return records, len(headers), cells


def _merged_fills(
    worksheet: Worksheet,
    max_columns: int,
    max_cells: int,
) -> dict[tuple[int, int], object]:
    """合并单元格的其余位置沿用左上角的值。"""
    fills: dict[tuple[int, int], object] = {}
    covered = 0
    for merged in worksheet.merged_cells.ranges:
        if merged.max_col > max_columns:
            raise ParseFailed(f"Excel 超过 {max_columns} 列")
        area = (merged.max_row - merged.min_row + 1) * (
            merged.max_col - merged.min_col + 1
        )
        covered += area
        if covered > max_cells:
            raise ParseFailed(f"Excel 超过 {max_cells} 个单元格")
        value = worksheet.cell(merged.min_row, merged.min_col).value
        for row in range(merged.min_row, merged.max_row + 1):
            for col in range(merged.min_col, merged.max_col + 1):
                if (row, col) != (merged.min_row, merged.min_col):
                    fills[(row, col)] = value
    return fills


def _materialize(
    values: list[object],
    row_number: int,
    fills: dict[tuple[int, int], object],
) -> list[object]:
    """把合并进来的值填进当前行。"""
    extra = [col for row, col in fills if row == row_number]
    if extra:
        width = max(extra)
        if len(values) < width:
            values.extend([None] * (width - len(values)))
    for index in range(len(values)):
        key = (row_number, index + 1)
        if values[index] is None and key in fills:
            values[index] = fills[key]
    while values and _value(values[-1]) is None:
        values.pop()
    return values


def _headers(values: list[object], used: set[str]) -> list[str]:
    """空表头和重复表头改成 column_N。"""
    headers: list[str] = []
    for index, value in enumerate(values, start=1):
        name = _header_text(value)
        if not name or name in used:
            name = _column_name(index, used)
        used.add(name)
        headers.append(name)
    return headers


def _extend_headers(
    headers: list[str],
    used: set[str],
    width: int,
    records: list[dict[str, object]],
) -> None:
    """数据行比表头更宽时补列表头，已有记录补 null。"""
    while len(headers) < width:
        name = _column_name(len(headers) + 1, used)
        used.add(name)
        headers.append(name)
        for record in records:
            record[name] = None


def _column_name(index: int, used: set[str]) -> str:
    """生成不与现有表头冲突的列名。"""
    name = f"column_{index}"
    while name in used:
        name = f"{name}_"
    return name


def _header_text(value: object) -> str:
    """表头统一成字符串。"""
    normalized = _value(value)
    if normalized is None:
        return ""
    return str(normalized).strip()


def _value(value: object) -> object:
    """保留数字和布尔值，日期转为 ISO 8601，空单元格为 null。"""
    if value is None:
        return None
    if isinstance(value, bool):
        return value
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%dT%H:%M:%S")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        if not isfinite(value):
            return None
        if value.is_integer():
            return int(value)
        return value
    if isinstance(value, str):
        text = value.strip()
        return text or None
    return str(value)


def _blank(values: list[object]) -> bool:
    """判断一行是否没有内容。"""
    return all(_value(cell) is None for cell in values)


def _check_cells(cells: int, max_cells: int) -> None:
    """单元格数量超过上限时拒绝。"""
    if cells > max_cells:
        raise ParseFailed(f"Excel 超过 {max_cells} 个单元格")


def _title(filename: str) -> str:
    """用文件名去掉扩展名作为标题。"""
    name = filename.rsplit("/", 1)[-1]
    stem, dot, _extension = name.rpartition(".")
    return stem if dot else name
