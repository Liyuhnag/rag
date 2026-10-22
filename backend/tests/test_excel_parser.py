# -*- coding: utf-8 -*-
"""Excel 按工作表转成 JSON。"""

import io
import json
from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import Workbook

from services.parsing.errors import ParseFailed
from services.parsing.models import ParseInput
from services.parsing.models import ParseRequest
from services.parsing.parsers.xls import XlsParser
from services.parsing.parsers.xlsx import XlsxParser


class _Reader:
    def __init__(self, content: bytes) -> None:
        self.content = content

    def read_object(self, key: str) -> bytes:
        return self.content

    def write_object(self, key: str, content: bytes, content_type: str) -> None:
        raise AssertionError("Excel 解析不应写入对象")


def _source(file_type: str = "xlsx") -> ParseInput:
    return ParseInput(
        request=ParseRequest(
            document_id="doc-1",
            version=1,
            object_key=f"documents/doc-1/销售.{file_type}",
            file_type=file_type,
            filename=f"销售.{file_type}",
            size_bytes=32,
        ),
    )


def _save(workbook: Workbook) -> bytes:
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _sales_book() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "销售"
    sheet["A1"] = "区域"
    sheet["B1"] = "金额"
    sheet["C1"] = "启用"
    sheet["D1"] = "日期"
    sheet.merge_cells("A2:A3")
    sheet["A2"] = "华东"
    sheet["B2"] = 100
    sheet["B3"] = 12.5
    sheet["C2"] = True
    sheet["D2"] = datetime(2026, 1, 2, 3, 4, 5)
    sheet["A4"] = "日期"
    sheet["B4"] = ""
    sheet["C4"] = "日期"
    sheet["D4"] = 1
    stock = workbook.create_sheet("库存")
    stock["A1"] = "商品"
    stock["A2"] = "铅笔"
    stock["B1"] = "金额"
    stock["B2"] = "=1+1"
    return _save(workbook)


def _records(block_text: str) -> list[dict[str, object]]:
    payload = json.loads(block_text)
    assert isinstance(payload, list)
    return payload


def test_xlsx_converts_each_sheet_to_records():
    parsed = XlsxParser(_Reader(_sales_book())).parse(_source())

    assert parsed.title == "销售"
    assert [block.sheet for block in parsed.blocks] == ["销售", "库存"]
    assert _records(parsed.blocks[0].text) == [
        {
            "区域": "华东",
            "金额": 100,
            "启用": True,
            "日期": "2026-01-02T03:04:05",
        },
        {"区域": "华东", "金额": 12.5, "启用": None, "日期": None},
        {"区域": "日期", "金额": None, "启用": "日期", "日期": 1},
    ]
    assert _records(parsed.blocks[1].text) == [{"商品": "铅笔", "金额": None}]


def test_xlsx_rejects_too_many_rows():
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet["A1"] = "金额"
    sheet["A2"] = 1
    sheet["A3"] = 2

    with pytest.raises(ParseFailed, match="超过 1 行"):
        XlsxParser(_Reader(_save(workbook)), max_rows=1).parse(_source())


def test_xlsx_rejects_too_many_sheets():
    workbook = Workbook()
    workbook.create_sheet("第二张")

    with pytest.raises(ParseFailed, match="超过 1 个工作表"):
        XlsxParser(_Reader(_save(workbook)), max_sheets=1).parse(_source())


def test_xlsx_rejects_invalid_file():
    with pytest.raises(ParseFailed, match="无法解析"):
        XlsxParser(_Reader(b"not-xlsx")).parse(_source())


def test_xls_converts_before_parsing():
    content = _sales_book()

    class _Converter:
        def to_xlsx(self, source_path: Path) -> Path:
            assert source_path.read_bytes() == b"old-xls"
            target = source_path.with_suffix(".xlsx")
            target.write_bytes(content)
            return target

    parsed = XlsParser(
        converter=_Converter(),
        reader=_Reader(b"old-xls"),
    ).parse(_source("xls"))

    assert parsed.blocks[0].sheet == "销售"
    assert _records(parsed.blocks[0].text)[0]["区域"] == "华东"
