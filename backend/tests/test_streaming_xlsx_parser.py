# -*- coding: utf-8 -*-
"""XLSX disk-backed streaming contracts."""

import io
import json
import struct
from datetime import datetime
from pathlib import Path
from zipfile import ZipFile
from zipfile import ZIP_DEFLATED

import pytest
from openpyxl import Workbook

from services.parsing.errors import ParseFailed
from services.parsing.models import ParseInput
from services.parsing.models import ParseRequest
from services.parsing.parsers.xlsx import XlsxParser


class DownloadReader:
    def __init__(self, content: bytes) -> None:
        self.content = content
        self.path = None

    def read_object(self, key: str) -> bytes:
        raise AssertionError("streaming must not read full object")

    def download_object(self, key: str, path: Path) -> None:
        self.path = path
        path.write_bytes(self.content)


def source() -> ParseInput:
    return ParseInput(ParseRequest("doc", 1, "book.xlsx", "xlsx", "book.xlsx", 1))


def save(book: Workbook) -> bytes:
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def test_streaming_preserves_merges_dates_and_late_columns():
    book = Workbook()
    sheet = book.active
    sheet.append(["name", "name"])
    sheet.append([" Alice ", 1])
    sheet.append([None, datetime(2026, 1, 2, 3, 4, 5), True])
    sheet.merge_cells("A2:A3")
    reader = DownloadReader(save(book))
    blocks = list(XlsxParser(reader, batch_rows=1).iter_blocks(source()))
    assert [json.loads(block.text) for block in blocks] == [
        [{"name": "Alice", "column_2": 1, "column_3": None}],
        [{"name": "Alice", "column_2": "2026-01-02T03:04:05", "column_3": True}],
    ]
    assert all(block.metadata["row_count"] == 1 for block in blocks)
    assert all(block.metadata["column_count"] == 3 for block in blocks)
    assert reader.path is not None and not reader.path.exists()


def test_streaming_reads_shared_strings_and_cached_formula():
    book = Workbook()
    sheet = book.active
    sheet.append(["title"])
    sheet.append(["unused"])
    original = save(book)
    output = io.BytesIO()
    with (
        ZipFile(io.BytesIO(original)) as archive,
        ZipFile(output, "w", ZIP_DEFLATED) as target,
    ):
        for item in archive.infolist():
            data = archive.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = (
                    b'<worksheet xmlns="http://schemas.openxmlformats.org/'
                    b'spreadsheetml/2006/main"><sheetData>'
                    b'<row r="1"><c r="A1" t="s"><v>0</v></c></row>'
                    b'<row r="2"><c r="A2" t="s"><v>1</v></c></row>'
                    b'<row r="3"><c r="A3"><f>1+1</f><v>2</v></c></row>'
                    b"</sheetData></worksheet>"
                )
            target.writestr(item.filename, data)
        target.writestr(
            "xl/sharedStrings.xml",
            (
                '<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
                "<si><t>title</t></si><si><r><t>shared</t></r>"
                "<r><t> string</t></r></si></sst>"
            ),
        )
    blocks = list(XlsxParser(DownloadReader(output.getvalue())).iter_blocks(source()))
    assert json.loads(blocks[0].text) == [{"title": "shared string"}, {"title": 2}]


def test_streaming_closing_generator_cleans_temporary_files():
    book = Workbook()
    book.active.append(["header"])
    book.active.append([1])
    book.active.append([2])
    reader = DownloadReader(save(book))
    stream = XlsxParser(reader, batch_rows=1).iter_blocks(source())
    next(stream)
    assert reader.path.exists()
    stream.close()
    assert not reader.path.exists()


@pytest.mark.parametrize(
    "options,match",
    [
        ({"max_rows": 1}, "超过 1 行"),
        ({"max_columns": 1}, "超过 1 列"),
        ({"max_cells": 2}, "超过 2 个单元格"),
        ({"max_uncompressed": 10}, "解压后超过"),
    ],
)
def test_streaming_enforces_existing_limits(options, match):
    book = Workbook()
    book.active.append(["a", "b"])
    book.active.append([1, 2])
    book.active.append([3, 4])
    with pytest.raises(ParseFailed, match=match):
        list(XlsxParser(DownloadReader(save(book)), **options).iter_blocks(source()))


def test_streaming_rejects_invalid_file_and_empty_workbook():
    with pytest.raises(ParseFailed, match="无法解析"):
        list(XlsxParser(DownloadReader(b"bad")).iter_blocks(source()))
    with pytest.raises(ParseFailed, match="没有表头"):
        list(XlsxParser(DownloadReader(save(Workbook()))).iter_blocks(source()))


def replace_member(content, name, payload):
    output = io.BytesIO()
    with (
        ZipFile(io.BytesIO(content)) as archive,
        ZipFile(output, "w", ZIP_DEFLATED) as target,
    ):
        for item in archive.infolist():
            target.writestr(
                item.filename,
                payload if item.filename == name else archive.read(item.filename),
            )
    return output.getvalue()


def test_streaming_rejects_oversized_xml_and_cleans_files():
    content = replace_member(
        save(Workbook()),
        "xl/worksheets/sheet1.xml",
        (
            b'<worksheet><sheetData><row r="1"><c r="A1" t="inlineStr"><is><t>'
            + b"x" * (1024 * 1024 + 1)
            + b"</t></is></c></row></sheetData></worksheet>"
        ),
    )
    reader = DownloadReader(content)
    with pytest.raises(ParseFailed, match="大小限制"):
        list(XlsxParser(reader).iter_blocks(source()))
    assert not reader.path.exists()


def test_streaming_preserves_wide_merge_and_empty_header_sheet():
    book = Workbook()
    book.active.append(["only header"])
    sheet = book.create_sheet("merged")
    sheet.append(["a", "b", "c"])
    sheet.append(["same"])
    sheet.merge_cells("A2:C3")
    blocks = list(XlsxParser(DownloadReader(save(book))).iter_blocks(source()))
    assert json.loads(blocks[0].text) == []
    assert json.loads(blocks[1].text) == [dict.fromkeys(["a", "b", "c"], "same")] * 2


def test_streaming_rejects_too_many_sheets():
    book = Workbook()
    book.create_sheet("extra")
    with pytest.raises(ParseFailed, match="超过 1 个工作表"):
        list(XlsxParser(DownloadReader(save(book)), max_sheets=1).iter_blocks(source()))


def test_streaming_rejects_xml_entity_declarations():
    payload = (
        b'<!DOCTYPE worksheet [<!ENTITY text "header">]>'
        b'<worksheet><sheetData><row r="1"><c r="A1" t="inlineStr">'
        b"<is><t>&text;</t></is></c></row></sheetData></worksheet>"
    )
    content = replace_member(save(Workbook()), "xl/worksheets/sheet1.xml", payload)
    with pytest.raises(ParseFailed, match="XML 声明"):
        list(XlsxParser(DownloadReader(content)).iter_blocks(source()))


@pytest.mark.parametrize("shared", [False, True])
def test_streaming_rich_text_excludes_phonetic_annotation(shared):
    rich = (
        '<r><t>東京</t></r><r><t>駅</t></r><rPh sb="0" eb="2"><t>とうきょう</t></rPh>'
    )
    cell = (
        '<c r="A2" t="s"><v>0</v></c>'
        if shared
        else ('<c r="A2" t="inlineStr"><is>' + rich + "</is></c>")
    )
    xml = (
        '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        '<sheetData><row r="1"><c r="A1" t="inlineStr">'
        '<is><t>place</t></is></c></row><row r="2">'
        + cell
        + "</row></sheetData></worksheet>"
    )
    content = replace_member(save(Workbook()), "xl/worksheets/sheet1.xml", xml)
    if shared:
        output = io.BytesIO()
        with (
            ZipFile(io.BytesIO(content)) as archive,
            ZipFile(output, "w", ZIP_DEFLATED) as target,
        ):
            for item in archive.infolist():
                data = archive.read(item.filename)
                if item.filename == "[Content_Types].xml":
                    data = data.replace(
                        b"</Types>",
                        b'<Override PartName="/xl/sharedStrings.xml" '
                        b'ContentType="application/vnd.openxmlformats-officedocument.'
                        b'spreadsheetml.sharedStrings+xml"/></Types>',
                    )
                target.writestr(item.filename, data)
            target.writestr(
                "xl/sharedStrings.xml",
                (
                    '<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
                    "<si>" + rich + "</si></sst>"
                ),
            )
        content = output.getvalue()
    parser = XlsxParser(DownloadReader(content))
    expected = json.loads(parser.parse_bytes(content, source()).blocks[0].text)
    actual = json.loads(list(parser.iter_blocks(source()))[0].text)
    assert actual == expected == [{"place": "東京駅"}]


def test_streaming_missing_cell_reference_follows_previous_column():
    xml = (
        '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        '<sheetData><row r="1">'
        '<c r="A1" t="inlineStr"><is><t>a</t></is></c>'
        '<c r="B1" t="inlineStr"><is><t>b</t></is></c>'
        '<c r="C1" t="inlineStr"><is><t>c</t></is></c></row>'
        '<row r="2"><c r="B2"><v>1</v></c><c><v>2</v></c></row>'
        "</sheetData></worksheet>"
    )
    content = replace_member(save(Workbook()), "xl/worksheets/sheet1.xml", xml)
    parser = XlsxParser(DownloadReader(content))
    expected = json.loads(parser.parse_bytes(content, source()).blocks[0].text)
    actual = json.loads(list(parser.iter_blocks(source()))[0].text)
    assert actual == expected == [{"a": None, "b": 1, "c": 2}]


@pytest.mark.parametrize("count,size", [(10001, 0), (1, 16 * 1024 * 1024 + 1)])
def test_streaming_rejects_directory_before_zipfile(monkeypatch, count, size):
    from services.parsing.parsers import _xlsx_stream

    content = struct.pack("<4s4H2LH", b"PK\x05\x06", 0, 0, count, count, size, 0, 0)

    def forbidden(*args, **kwargs):
        raise AssertionError("ZipFile must not allocate directory before preflight")

    monkeypatch.setattr(_xlsx_stream, "ZipFile", forbidden)
    with pytest.raises(ParseFailed, match="Excel.*目录"):
        list(XlsxParser(DownloadReader(content)).iter_blocks(source()))
